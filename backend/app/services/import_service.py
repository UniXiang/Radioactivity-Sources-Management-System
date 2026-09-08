from __future__ import annotations

import hashlib
import json
import shutil
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.database.session import Database
from app.repositories.source_json_repository import SourceJsonRepository
from app.schemas.import_xlsx import ImportMapping
from app.services.audit_service import AuditService, now_iso
from app.services.calibration_service import CalibrationService
from app.services.activity_service import ActivityService


class ImportService:
    def __init__(self, database: Database, sources: SourceJsonRepository, activity: ActivityService, directory: str | Path):
        self.database = database
        self.sources = sources
        self.activity = activity
        self.directory = Path(directory)
        self.directory.mkdir(parents=True, exist_ok=True)
        self.calibrations = CalibrationService(database, sources, activity)
        self.audit = AuditService(database)

    def save_upload(self, filename: str, content: bytes) -> dict[str, Any]:
        digest = hashlib.sha256(content).hexdigest()
        duplicate = self.database.fetch_one("SELECT id, filename, uploaded_at FROM import_batches WHERE sha256 = ? LIMIT 1", (digest,))
        token = uuid.uuid4().hex
        path = self.directory / f"{token}.xlsx"
        path.write_bytes(content)
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
            preview = [_serialize_row(row) for _, row in zip(range(20), rows)]
            sheets.append({"name": sheet_name, "headers": headers, "preview": preview, "max_row": sheet.max_row, "max_column": sheet.max_column})
        workbook.close()
        return {"upload_token": token, "filename": filename, "size": len(content), "sha256": digest, "duplicate": dict(duplicate) if duplicate else None, "sheets": sheets}

    def validate(self, upload_token: str, sheet_name: str, mapping: ImportMapping) -> dict[str, Any]:
        path = self._path(upload_token)
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"sheet {sheet_name} not found")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
        rows = []
        last_date = None
        for number, values in enumerate(iterator, start=2):
            raw = {headers[index] if index < len(headers) else f"Column {index + 1}": _serialize(value) for index, value in enumerate(values)}
            if not any(value not in (None, "") for value in raw.values()):
                continue
            normalized, messages = self._normalize(raw, mapping)
            if mapping.calibration_date:
                if normalized.get("calibration_date"):
                    last_date = normalized["calibration_date"]
                elif last_date:
                    normalized["calibration_date"] = last_date
                    messages.append({"severity": "warning", "code": "DATE_CARRIED_FORWARD", "message": "Blank date inherited from the previous non-empty date"})
            if normalized.get("source_label_raw"):
                normalized["source_id"] = self._match_source(normalized["source_label_raw"])
                if normalized["source_id"] is None:
                    messages.append({"severity": "warning", "code": "SOURCE_UNMATCHED", "message": "Source label did not match a Registry source"})
            if not normalized.get("calibration_date") and not normalized.get("start_time"):
                severity = "warning" if not mapping.calibration_date else "warning"
                messages.append({"severity": severity, "code": "DATE_MISSING", "message": "No calibration date was mapped for this row"})
            rows.append({"row": number, "raw_data": raw, "normalized": normalized, "messages": messages})
        workbook.close()
        seen = set()
        for row in rows:
            identity = json.dumps({key: row["normalized"].get(key) for key in ("source_id", "source_label_raw", "calibration_date", "run_number", "position_label")}, ensure_ascii=False, sort_keys=True, default=str)
            if identity in seen:
                row["messages"].append({"severity": "warning", "code": "DUPLICATE_ROW", "message": "Duplicate normalized row; it will be skipped by default"})
            seen.add(identity)
        counts = {
            "total_rows": len(rows),
            "warning_rows": sum(any(message["severity"] == "warning" for message in row["messages"]) for row in rows),
            "failed_rows": sum(any(message["severity"] == "error" for message in row["messages"]) for row in rows),
            "unmatched_sources": sum(any(message["code"] == "SOURCE_UNMATCHED" for message in row["messages"]) for row in rows),
            "duplicate_rows": sum(any(message["code"] == "DUPLICATE_ROW" for message in row["messages"]) for row in rows),
        }
        return {"sheet_name": sheet_name, "headers": headers, "rows": rows, "summary": counts}

    def confirm(self, *, upload_token: str, filename: str, sha256: str, sheet_name: str, mapping: ImportMapping, user: dict | None, ip_address: str | None = None, duplicate_policy: str = "skip") -> dict[str, Any]:
        existing = self.database.fetch_one("SELECT id FROM import_batches WHERE sha256 = ? LIMIT 1", (sha256,))
        if existing and duplicate_policy not in {"allow", "reimport"}:
            raise ValueError("该文件已经导入。相同文件的重复导入已阻止。")
        validation = self.validate(upload_token, sheet_name, mapping)
        rows = validation["rows"]
        with self.database.transaction() as connection:
            cursor = connection.execute(
                """INSERT INTO import_batches (filename, sha256, sheet_name, total_rows, uploaded_by, uploaded_at, status)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (filename, sha256, sheet_name, validation["summary"]["total_rows"], _id(user), now_iso(), "processing"),
            )
            batch_id = int(cursor.lastrowid)
            imported = warnings = failed = 0
            identities = set()
            for row in rows:
                messages = row["messages"]
                if any(message["severity"] == "error" for message in messages):
                    failed += 1
                    continue
                identity = json.dumps({key: row["normalized"].get(key) for key in ("source_id", "source_label_raw", "calibration_date", "run_number", "position_label")}, ensure_ascii=False, sort_keys=True, default=str)
                if identity in identities and duplicate_policy == "skip":
                    warnings += 1
                    continue
                identities.add(identity)
                if any(message["severity"] == "warning" for message in messages):
                    warnings += 1
                self.calibrations.create(row["normalized"], user, ip_address, connection=connection, import_batch_id=batch_id, original_sheet=sheet_name, original_row=row["row"], raw_data=row["raw_data"])
                imported += 1
            status = "completed" if failed == 0 else "completed_with_errors"
            connection.execute("UPDATE import_batches SET imported_rows = ?, warning_rows = ?, failed_rows = ?, status = ? WHERE id = ?", (imported, warnings, failed, status, batch_id))
            self.audit.record(user_id=_id(user), username=_name(user), action="xlsx_import", entity_type="import_batch", entity_id=str(batch_id), new_data={"filename": filename, "sheet_name": sheet_name, "imported_rows": imported, "failed_rows": failed}, ip_address=ip_address, connection=connection)
        return dict(self.database.fetch_one("SELECT * FROM import_batches WHERE id = ?", (batch_id,)))

    def list_batches(self) -> list[dict[str, Any]]:
        return [dict(row) for row in self.database.fetch_all("SELECT * FROM import_batches ORDER BY uploaded_at DESC, id DESC")]

    def _path(self, token: str) -> Path:
        path = self.directory / f"{token}.xlsx"
        if not path.exists() or path.name != f"{token}.xlsx":
            raise ValueError("upload token not found")
        return path

    def _match_source(self, label: str) -> str | None:
        target = _canonical(label)
        for source in self.sources.list_sources():
            if target in {_canonical(source.source_id), _canonical(source.display_name), _canonical(source.isotope), _canonical(source.serial_number)}:
                return source.source_id
        return None

    def _normalize(self, raw: dict[str, Any], mapping: ImportMapping) -> tuple[dict[str, Any], list[dict[str, str]]]:
        result: dict[str, Any] = {}
        messages: list[dict[str, str]] = []
        for field, column in mapping.model_dump().items():
            if not column:
                continue
            value = raw.get(column)
            if field in {"calibration_date", "start_time", "end_time"} and value not in (None, ""):
                try:
                    result[field] = _parse_date_value(value).isoformat() if field == "calibration_date" else _parse_date_value(value).isoformat()
                except ValueError:
                    messages.append({"severity": "error", "code": "DATE_INVALID", "message": f"Cannot parse {field}: {value}"})
                    continue
            elif value not in (None, ""):
                result[field] = str(value).strip() if field in {"run_number", "position_label", "operator", "purpose", "comment", "source_label_raw", "isotope", "calibration_system"} else value
        return result, messages


def _serialize(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _serialize_row(row: tuple[Any, ...]) -> list[Any]:
    return [_serialize(value) for value in row]


def _parse_date_value(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.utc)
    text = str(value).strip().replace("/", "-").replace("Z", "+00:00")
    if text.isdigit() and len(text) == 8:
        try:
            return datetime.strptime(text, "%Y%m%d").replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    for parser in (datetime.fromisoformat,):
        try:
            parsed = parser(text)
            return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    raise ValueError(text)


def _canonical(value: str) -> str:
    return "".join(character.lower() for character in str(value) if character.isalnum())


def _id(user: dict | None) -> int | None:
    return int(user["id"]) if user else None


def _name(user: dict | None) -> str | None:
    return user.get("username") if user else None
